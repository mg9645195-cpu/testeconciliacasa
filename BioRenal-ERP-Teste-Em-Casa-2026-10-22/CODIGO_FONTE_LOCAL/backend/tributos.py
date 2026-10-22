# -*- coding: utf-8 -*-
"""Controle local de notas e estimativas tributárias; não substitui a apuração fiscal."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import json
import re
import sqlite3

from openpyxl import load_workbook

from database import get_conn

CENT = Decimal("0.01")
MONTHS = {"JAN": 1, "FEV": 2, "MAR": 3, "ABR": 4, "MAI": 5, "JUN": 6,
          "JUL": 7, "AGO": 8, "SET": 9, "OUT": 10, "NOV": 11, "DEZ": 12}
DEFAULT_CONFIG = {"pis": "0.0065", "cofins": "0.03", "irpj_presuncao": "",
                  "irpj_aliquota": "", "irpj_adicional": "", "irpj_limite_trimestre": "",
                  "csll_presuncao": "", "csll_aliquota": ""}


def cents(value) -> int:
    return int((Decimal(str(value or 0)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def money(value: int) -> float:
    return float((Decimal(value) / 100).quantize(CENT))


def amount(base_cents: int, rate: str) -> int:
    return int((Decimal(base_cents) * Decimal(str(rate))).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _date(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    result = str(value)[:10]
    date.fromisoformat(result)
    return result


def _competence(description: str, emission: str) -> str:
    match = re.search(r"REF\.?\s*A?\s*(JAN|FEV|MAR|ABR|MAI|JUN|JUL|AGO|SET|OUT|NOV|DEZ)\s*/\s*(20\d\d)", description.upper())
    return f"{match.group(2)}-{MONTHS[match.group(1)]:02d}" if match else emission[:7]


def ensure_tax_schema(conn: sqlite3.Connection) -> None:
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS tributos_notas (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          tipo TEXT NOT NULL CHECK(tipo IN ('EMITIDA','TOMADA')),
          numero TEXT NOT NULL,
          empresa TEXT NOT NULL,
          emissao TEXT NOT NULL,
          competencia_servico TEXT NOT NULL,
          data_financeira TEXT NOT NULL DEFAULT '',
          descricao TEXT NOT NULL DEFAULT '',
          bruto_centavos INTEGER NOT NULL CHECK(bruto_centavos >= 0),
          ret_pis_centavos INTEGER NOT NULL DEFAULT 0,
          ret_cofins_centavos INTEGER NOT NULL DEFAULT 0,
          ret_csll_centavos INTEGER NOT NULL DEFAULT 0,
          ret_irrf_centavos INTEGER NOT NULL DEFAULT 0,
          origem TEXT NOT NULL DEFAULT 'MANUAL',
          criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
          atualizado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
          UNIQUE(tipo, numero, empresa, emissao)
        );
        CREATE TABLE IF NOT EXISTS tributos_config (
          chave TEXT PRIMARY KEY, valor TEXT NOT NULL, atualizado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS tributos_auditoria (
          id INTEGER PRIMARY KEY AUTOINCREMENT, nota_id INTEGER, acao TEXT NOT NULL,
          antes_json TEXT, depois_json TEXT, criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS tributos_guias (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          tributo TEXT NOT NULL, competencia TEXT NOT NULL,
          tipo TEXT NOT NULL CHECK(tipo IN ('MENSAL','TRIMESTRAL','AVULSO')),
          documento TEXT NOT NULL DEFAULT '',
          base_centavos INTEGER NOT NULL DEFAULT 0,
          aliquota TEXT NOT NULL DEFAULT '',
          valor_centavos INTEGER NOT NULL CHECK(valor_centavos >= 0),
          mes_previsto_pagamento TEXT NOT NULL,
          data_pagamento TEXT NOT NULL DEFAULT '',
          observacao TEXT NOT NULL DEFAULT '',
          atualizado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );
        CREATE INDEX IF NOT EXISTS idx_tributos_recebimento ON tributos_notas(tipo, data_financeira);
    """)
    for key, value in DEFAULT_CONFIG.items():
        conn.execute("INSERT OR IGNORE INTO tributos_config(chave, valor) VALUES (?, ?)", (key, value))
    conn.commit()


def _config(conn):
    return {r["chave"]: r["valor"] for r in conn.execute("SELECT chave, valor FROM tributos_config")}


def _note_dict(row):
    item = dict(row)
    for field in ("bruto", "ret_pis", "ret_cofins", "ret_csll", "ret_irrf"):
        item[field] = money(item.pop(field + "_centavos"))
    item["liquido"] = money(cents(item["bruto"]) - sum(cents(item[f"ret_{name}"]) for name in ("pis", "cofins", "csll", "irrf")))
    return item


def save_note(data: dict) -> dict:
    kind = str(data.get("tipo", "")).upper()
    if kind not in ("EMITIDA", "TOMADA"):
        raise ValueError("Tipo de nota inválido.")
    number = str(data.get("numero", "")).strip()
    company = str(data.get("empresa", "")).strip()
    emission = _date(data.get("emissao"))
    financial = _date(data.get("data_financeira"))
    competence = str(data.get("competencia_servico") or emission[:7]).strip()
    if not number or not company or not emission or not re.fullmatch(r"20\d\d-(0[1-9]|1[0-2])", competence):
        raise ValueError("Informe número, empresa, emissão e competência (AAAA-MM).")
    gross = cents(data.get("bruto"))
    if gross <= 0:
        raise ValueError("Valor bruto precisa ser maior que zero.")
    withheld = [cents(data.get("ret_" + tax)) for tax in ("pis", "cofins", "csll", "irrf")]
    if any(v < 0 for v in withheld) or sum(withheld) > gross:
        raise ValueError("Retenções inválidas ou superiores ao valor bruto.")
    description = str(data.get("descricao", "")).strip()[:500]
    note_id = int(data.get("id") or 0)
    conn = get_conn()
    try:
        ensure_tax_schema(conn)
        before = conn.execute("SELECT * FROM tributos_notas WHERE id=?", (note_id,)).fetchone() if note_id else None
        if note_id and before is None:
            raise ValueError("Nota não encontrada.")
        fields = (kind, number, company, emission, competence, financial, description, gross, *withheld)
        if before:
            conn.execute("""UPDATE tributos_notas SET tipo=?,numero=?,empresa=?,emissao=?,competencia_servico=?,
                data_financeira=?,descricao=?,bruto_centavos=?,ret_pis_centavos=?,ret_cofins_centavos=?,
                ret_csll_centavos=?,ret_irrf_centavos=?,atualizado_em=CURRENT_TIMESTAMP WHERE id=?""", (*fields, note_id))
        else:
            note_id = conn.execute("""INSERT INTO tributos_notas(tipo,numero,empresa,emissao,competencia_servico,
                data_financeira,descricao,bruto_centavos,ret_pis_centavos,ret_cofins_centavos,
                ret_csll_centavos,ret_irrf_centavos) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""", fields).lastrowid
        after = conn.execute("SELECT * FROM tributos_notas WHERE id=?", (note_id,)).fetchone()
        conn.execute("INSERT INTO tributos_auditoria(nota_id,acao,antes_json,depois_json) VALUES(?,?,?,?)",
                     (note_id, "EDITAR" if before else "CRIAR", json.dumps(dict(before), ensure_ascii=False) if before else None,
                      json.dumps(dict(after), ensure_ascii=False)))
        conn.commit()
        return _note_dict(after)
    except sqlite3.IntegrityError as exc:
        conn.rollback()
        raise ValueError("Esta nota já existe para empresa e data de emissão.") from exc
    finally:
        conn.close()


def save_config(data: dict) -> dict:
    values = data.get("config") or {}
    if not isinstance(values, dict):
        raise ValueError("Configuração inválida.")
    conn = get_conn()
    try:
        ensure_tax_schema(conn)
        for key, raw in values.items():
            if key not in DEFAULT_CONFIG:
                raise ValueError("Parâmetro tributário desconhecido.")
            value = str(raw).strip().replace(",", ".")
            if value:
                number = Decimal(value)
                if number < 0 or (key != "irpj_limite_trimestre" and number > 1):
                    raise ValueError(f"Parâmetro {key} fora do intervalo permitido.")
            conn.execute("UPDATE tributos_config SET valor=?,atualizado_em=CURRENT_TIMESTAMP WHERE chave=?", (value, key))
        conn.commit()
        return _config(conn)
    finally:
        conn.close()


def _plus_month(period: str) -> str:
    year, month = map(int, period.split("-"))
    return f"{year + (month == 12)}-{(month % 12) + 1:02d}"


def save_tax_guide(data: dict) -> dict:
    tax = str(data.get("tributo", "")).strip().upper()[:80]
    period = str(data.get("competencia", "")).strip()
    kind = str(data.get("tipo", "AVULSO")).strip().upper()
    predicted = str(data.get("mes_previsto_pagamento", "")).strip()
    paid = _date(data.get("data_pagamento"))
    for value in (period, predicted):
        if not re.fullmatch(r"20\d\d-(0[1-9]|1[0-2])", value):
            raise ValueError("Competência e mês previsto devem estar no formato AAAA-MM.")
    if not tax or kind not in ("MENSAL", "TRIMESTRAL", "AVULSO"):
        raise ValueError("Informe tributo e periodicidade válidos.")
    base = cents(data.get("base"))
    rate = str(data.get("aliquota") or "").strip().replace(",", ".")
    if rate and (Decimal(rate) < 0 or Decimal(rate) > 1):
        raise ValueError("Alíquota deve ficar entre 0 e 100%.")
    value = amount(base, rate) if rate and base else cents(data.get("valor"))
    if value < 0:
        raise ValueError("Valor não pode ser negativo.")
    guide_id = int(data.get("id") or 0)
    fields = (tax, period, kind, str(data.get("documento", "")).strip()[:80], base, rate,
              value, predicted, paid, str(data.get("observacao", "")).strip()[:500])
    conn = get_conn()
    try:
        ensure_tax_schema(conn)
        before = conn.execute("SELECT * FROM tributos_guias WHERE id=?", (guide_id,)).fetchone() if guide_id else None
        if guide_id and not before:
            raise ValueError("Guia não encontrada.")
        if before:
            conn.execute("""UPDATE tributos_guias SET tributo=?,competencia=?,tipo=?,documento=?,base_centavos=?,
                aliquota=?,valor_centavos=?,mes_previsto_pagamento=?,data_pagamento=?,observacao=?,
                atualizado_em=CURRENT_TIMESTAMP WHERE id=?""", (*fields, guide_id))
        else:
            guide_id = conn.execute("""INSERT INTO tributos_guias(tributo,competencia,tipo,documento,base_centavos,
                aliquota,valor_centavos,mes_previsto_pagamento,data_pagamento,observacao)
                VALUES(?,?,?,?,?,?,?,?,?,?)""", fields).lastrowid
        after = conn.execute("SELECT * FROM tributos_guias WHERE id=?", (guide_id,)).fetchone()
        conn.execute("INSERT INTO tributos_auditoria(nota_id,acao,antes_json,depois_json) VALUES(NULL,?,?,?)",
                     ("EDITAR_GUIA" if before else "CRIAR_GUIA", json.dumps(dict(before), ensure_ascii=False) if before else None,
                      json.dumps(dict(after), ensure_ascii=False)))
        conn.commit()
        return _guide_dict(after)
    finally:
        conn.close()


def _guide_dict(row):
    result = dict(row)
    result["base"] = money(result.pop("base_centavos"))
    result["valor"] = money(result.pop("valor_centavos"))
    result["status"] = "PAGO" if result["data_pagamento"] else "PREVISTO"
    return result


def get_tributos(year: int = 2026) -> dict:
    if not 2020 <= year <= 2100:
        raise ValueError("Ano inválido.")
    conn = get_conn()
    try:
        ensure_tax_schema(conn)
        notes = [_note_dict(r) for r in conn.execute("SELECT * FROM tributos_notas WHERE substr(emissao,1,4)=? OR substr(data_financeira,1,4)=? ORDER BY emissao,id", (str(year), str(year)))]
        guides = [_guide_dict(r) for r in conn.execute("SELECT * FROM tributos_guias WHERE substr(competencia,1,4)=? OR substr(data_pagamento,1,4)=? ORDER BY competencia,id", (str(year), str(year)))]
        config = _config(conn)
    finally:
        conn.close()
    monthly = defaultdict(lambda: {"bruto_centavos": 0, "ret_irrf_centavos": 0, "notas": 0})
    for note in notes:
        if note["tipo"] != "EMITIDA" or not note["data_financeira"]:
            continue
        period = note["data_financeira"][:7]
        monthly[period]["bruto_centavos"] += cents(note["bruto"])
        monthly[period]["ret_irrf_centavos"] += cents(note["ret_irrf"])
        monthly[period]["notas"] += 1
    obligations = []
    quarters = defaultdict(int)
    for period, item in sorted(monthly.items()):
        gross = item["bruto_centavos"]
        if gross:
            for tax in ("pis", "cofins"):
                obligations.append({"tributo": tax.upper(), "periodo_apuracao": period,
                                    "mes_previsto_pagamento": _plus_month(period), "base": money(gross),
                                    "valor_estimado": money(amount(gross, config[tax])), "tipo": "MENSAL",
                                    "status": "ESTIMATIVA"})
            year_num, month_num = map(int, period.split("-"))
            quarters[(year_num, (month_num - 1) // 3 + 1)] += gross
    quarter_rows = []
    for (year_num, quarter), gross in sorted(quarters.items()):
        end_period = f"{year_num}-{quarter * 3:02d}"
        q = {"trimestre": f"{quarter}º/{year_num}", "base_recebida": money(gross),
             "mes_previsto_pagamento": _plus_month(end_period), "irpj_estimado": None, "csll_estimada": None}
        for tax in ("irpj", "csll"):
            pres = config.get(f"{tax}_presuncao")
            rate = config.get(f"{tax}_aliquota")
            if not pres or not rate:
                continue
            base = amount(gross, pres)
            due = amount(base, rate)
            if tax == "irpj" and config.get("irpj_adicional") and config.get("irpj_limite_trimestre"):
                due += amount(max(0, base - cents(config["irpj_limite_trimestre"])), config["irpj_adicional"])
            q["irpj_estimado" if tax == "irpj" else "csll_estimada"] = money(due)
            obligations.append({"tributo": tax.upper(), "periodo_apuracao": q["trimestre"],
                                "mes_previsto_pagamento": q["mes_previsto_pagamento"],
                                "base": money(base), "valor_estimado": money(due), "tipo": "TRIMESTRAL", "status": "ESTIMATIVA"})
        quarter_rows.append(q)
    summary = [{"mes": period, "notas_recebidas": item["notas"], "bruto_recebido": money(item["bruto_centavos"]),
                "irrf_retido": money(item["ret_irrf_centavos"])} for period, item in sorted(monthly.items())]
    return {"notas": notes, "guias": guides, "config": config, "mensal": summary, "trimestral": quarter_rows,
            "obrigacoes": sorted(obligations, key=lambda row: (row["mes_previsto_pagamento"], row["tributo"])),
            "sem_data_recebimento": sum(n["tipo"] == "EMITIDA" and not n["data_financeira"] for n in notes),
            "criterio": "Previsão operacional sobre notas emitidas com recebimento informado. Competência do serviço e mês de recebimento são preservados separadamente; PIS/Cofins são projetados para o mês seguinte ao recebimento, conforme regra informada pela clínica. Valores e vencimentos devem ser conferidos pela contabilidade."}


def import_tax_workbook(raw: bytes) -> dict:
    if len(raw) > 10_000_000:
        raise ValueError("Planilha acima do limite de 10 MB.")
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    if len(workbook.worksheets) < 2:
        raise ValueError("A planilha precisa conter as abas de notas emitidas e tomadas.")
    received = 0
    skipped = 0
    conn = get_conn()
    try:
        ensure_tax_schema(conn)
        for sheet_index, kind in ((0, "EMITIDA"), (1, "TOMADA")):
            ws = workbook.worksheets[sheet_index]
            for row in list(ws.iter_rows(min_row=3, values_only=True)):
                if kind == "EMITIDA":
                    number, company, emission, financial, description, gross = (row[i] if len(row) > i else None for i in range(6))
                    # Na aba emitida, a fórmula do líquido desconta somente IRRF (coluna J).
                    retention = (0, 0, 0, cents(row[9]) if len(row) > 9 else 0)
                else:
                    number, company, emission, gross = (row[i] if len(row) > i else None for i in (1, 2, 3, 4))
                    financial, description = None, ""
                    retention = tuple(cents(row[i]) if len(row) > i else 0 for i in (5, 6, 7, 8))
                if not number or not company or not emission or not gross:
                    continue
                emission_iso = _date(emission)
                financial_iso = _date(financial)
                competence = _competence(str(description or ""), emission_iso)
                values = (kind, str(number).strip(), str(company).strip(), emission_iso, competence,
                          financial_iso, str(description or ""), cents(gross), *retention)
                cursor = conn.execute("""INSERT OR IGNORE INTO tributos_notas(tipo,numero,empresa,emissao,competencia_servico,
                    data_financeira,descricao,bruto_centavos,ret_pis_centavos,ret_cofins_centavos,
                    ret_csll_centavos,ret_irrf_centavos,origem) VALUES(?,?,?,?,?,?,?,?,?,?,?,?, 'PLANILHA')""", values)
                received += cursor.rowcount
                skipped += not cursor.rowcount
        conn.commit()
        return {"importadas": received, "ja_existentes": skipped}
    finally:
        conn.close()
        workbook.close()
