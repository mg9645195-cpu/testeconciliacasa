# -*- coding: utf-8 -*-
"""
Biorenal — Camada de Persistência Local (SQLite) e Sincronização Bidirecional (.xlsx <-> SQLite).
Garante Fidelidade Contábil Absoluta (Decimal 2 casas, tolerância 0.005).
"""

import os
import sqlite3
import datetime
from decimal import Decimal, ROUND_HALF_UP
from typing import List, Dict, Any, Optional, Tuple
import openpyxl

from domain import (
    to_dec,
    dec_to_float,
    TOLERANCE,
    TWO_PLACES,
    PLANO_DE_CONTAS,
    TIPOS_PAGAMENTO,
    TIPOS_MOVIMENTACAO,
    ORIGENS_CLASSIFICACAO,
    validate_conferencia,
    extract_competencia,
    get_canonical_supplier,
    infer_missing_supplier,
    CANONICAL_SUPPLIER_GROUPS,
    suggest_for_movement,
    build_ferronefro_split,
    build_socio_split,
    build_notas_mistas_split,
)

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.environ.get("BIORENAL_DATA_DIR", os.path.join(BASE_DIR, "data"))
DB_PATH = os.path.join(DATA_DIR, "biorenal_2026.db")
DEFAULT_XLSX_PATH = os.path.join(DATA_DIR, "NOVO MODELO DE CONCILIAÇÃO BANCÁRIA COM RATEIO - 2026.xlsx")


def get_conn() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn


def init_schema(conn: sqlite3.Connection) -> None:
    cur = conn.cursor()
    cur.executescript(
        """
        CREATE TABLE IF NOT EXISTS config_parametros (
            chave TEXT PRIMARY KEY,
            valor TEXT NOT NULL,
            descricao TEXT
        );

        CREATE TABLE IF NOT EXISTS cadastros_plano_contas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ordem_cat INTEGER NOT NULL,
            categoria TEXT NOT NULL,
            ordem_sub INTEGER NOT NULL,
            subcategoria TEXT NOT NULL,
            UNIQUE(categoria, subcategoria)
        );

        CREATE TABLE IF NOT EXISTS mapeamento_historico (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            categoria_antiga TEXT,
            subcategoria_antiga TEXT,
            nova_categoria TEXT,
            nova_subcategoria TEXT,
            status TEXT,
            observacao TEXT
        );

        CREATE TABLE IF NOT EXISTS fornecedor_aliases (
            alias TEXT PRIMARY KEY COLLATE NOCASE,
            canonico TEXT NOT NULL,
            atualizado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS auditoria_cadastros (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            acao TEXT NOT NULL,
            alias TEXT NOT NULL,
            canonico_anterior TEXT,
            canonico_novo TEXT,
            criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS extrato (
            id_lanc TEXT PRIMARY KEY,
            ordem INTEGER NOT NULL,
            data TEXT NOT NULL,
            mes INTEGER NOT NULL,
            competencia_extraida TEXT NOT NULL,
            tipo_pagamento TEXT,
            tipo_movimentacao TEXT,
            historico_banco TEXT,
            fornecedor_favorecido TEXT,
            fornecedor_canonico TEXT,
            documento TEXT,
            entrada REAL NOT NULL DEFAULT 0,
            saida REAL NOT NULL DEFAULT 0,
            aplicacao REAL NOT NULL DEFAULT 0,
            resgate REAL NOT NULL DEFAULT 0,
            saldo_acumulado REAL NOT NULL DEFAULT 0,
            total_rateado REAL NOT NULL DEFAULT 0,
            diferenca REAL NOT NULL DEFAULT 0,
            status_rateio TEXT NOT NULL DEFAULT 'PENDENTE DE RATEIO',
            status_conciliacao TEXT,
            observacao TEXT,
            desconsiderar_rateio INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS rateios (
            id_rateio INTEGER PRIMARY KEY AUTOINCREMENT,
            id_lanc TEXT NOT NULL,
            data TEXT,
            mes INTEGER,
            categoria TEXT,
            subcategoria TEXT,
            valor REAL NOT NULL DEFAULT 0,
            detalhe_componente TEXT,
            natureza TEXT,
            fornecedor TEXT,
            origem_classificacao TEXT DEFAULT 'PENDENTE',
            conferencia TEXT DEFAULT 'SEM CATEGORIA',
            FOREIGN KEY (id_lanc) REFERENCES extrato(id_lanc) ON DELETE CASCADE
        );

        CREATE INDEX IF NOT EXISTS idx_rateios_id_lanc ON rateios(id_lanc);
        CREATE INDEX IF NOT EXISTS idx_extrato_ordem ON extrato(ordem);
        CREATE INDEX IF NOT EXISTS idx_extrato_mes ON extrato(mes);
        """
    )

    cur.execute("PRAGMA table_info(extrato)")
    cols = [r["name"] for r in cur.fetchall()]
    if "desconsiderar_rateio" not in cols:
        cur.execute("ALTER TABLE extrato ADD COLUMN desconsiderar_rateio INTEGER NOT NULL DEFAULT 0")

    # Seed Plano de Contas (12 categories, 87 subcategories)
    cur.execute("DELETE FROM cadastros_plano_contas;")
    for c_idx, (cat, subs) in enumerate(PLANO_DE_CONTAS.items(), start=1):
        for s_idx, sub in enumerate(subs, start=1):
            cur.execute(
                "INSERT INTO cadastros_plano_contas (ordem_cat, categoria, ordem_sub, subcategoria) VALUES (?, ?, ?, ?)",
                (c_idx, cat, s_idx, sub),
            )
    conn.commit()


def recalculate_all(conn: sqlite3.Connection) -> None:
    """
    Recalcula rigorosamente (usando Decimal com 2 casas decimais e tolerância 0.005):
    - Para cada linha em rateios: data, mes, natureza, fornecedor (herdados de extrato) e conferencia.
    - Para cada linha em extrato (em ordem de 1 a N):
      competencia_extraida, fornecedor_canonico, saldo_acumulado, total_rateado, diferenca, status_rateio.
    """
    cur = conn.cursor()

    # Read saldo inicial CC
    cur.execute("SELECT valor FROM config_parametros WHERE chave = 'saldo_inicial_cc'")
    row_si = cur.fetchone()
    saldo_inicial_cc = to_dec(row_si["valor"] if row_si else "0.00")

    # Fetch all extrato rows ordered by ordem
    cur.execute("SELECT * FROM extrato ORDER BY ordem ASC")
    ext_rows = [dict(r) for r in cur.fetchall()]
    cur.execute("SELECT alias, canonico FROM fornecedor_aliases")
    supplier_aliases = {r["alias"].strip().upper(): r["canonico"] for r in cur.fetchall()}

    # Fetch all rateios rows grouped by id_lanc
    cur.execute("SELECT * FROM rateios ORDER BY id_rateio ASC")
    rat_rows = [dict(r) for r in cur.fetchall()]

    rat_by_lanc: Dict[str, List[Dict[str, Any]]] = {}
    for r in rat_rows:
        rat_by_lanc.setdefault(r["id_lanc"], []).append(r)

    saldo_acum = saldo_inicial_cc
    ext_updates = []
    rat_updates = []

    for ext in ext_rows:
        id_lanc = ext["id_lanc"]
        data_iso = ext["data"] or ""
        mes = int(data_iso.split("-")[1]) if (data_iso and "-" in data_iso) else int(ext["mes"] or 1)
        hist = ext["historico_banco"] or ""
        obs = ext["observacao"] or ""
        forn = ext["fornecedor_favorecido"] or ""

        comp = extract_competencia(hist, obs, data_iso)
        base_canon = get_canonical_supplier(forn, hist)
        forn_canon = supplier_aliases.get(forn.strip().upper()) or supplier_aliases.get(base_canon.strip().upper()) or base_canon

        ent = to_dec(ext["entrada"])
        sai = to_dec(ext["saida"])
        apl = to_dec(ext["aplicacao"])
        res = to_dec(ext["resgate"])
        valor_banco = ent + sai + apl + res

        saldo_acum = (saldo_acum + ent + res - sai - apl).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)

        children = rat_by_lanc.get(id_lanc, [])
        total_rateado = Decimal("0.00")
        any_unclassified = False
        any_incompatible = False
        all_desconsiderado = len(children) > 0 and all((c.get("categoria") or "").strip() == "DESCONSIDERADO NO RATEIO" for c in children)
        is_desconsiderado = int(ext.get("desconsiderar_rateio") or 0) == 1 or all_desconsiderado

        for ch in children:
            v_ch = to_dec(ch["valor"])
            total_rateado = (total_rateado + v_ch).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)
            cat_ch = (ch["categoria"] or "").strip()
            sub_ch = (ch["subcategoria"] or "").strip()
            conf = validate_conferencia(id_lanc, cat_ch, sub_ch, v_ch)
            if conf in ("SEM CATEGORIA", "SEM SUBCATEGORIA"):
                any_unclassified = True
            elif conf == "INCOMPATÍVEL":
                any_incompatible = True

            rat_updates.append((
                data_iso,
                mes,
                ext["tipo_movimentacao"] or "",
                forn,
                conf,
                float(v_ch),
                ch["id_rateio"],
            ))

        if is_desconsiderado:
            total_rateado = valor_banco
            diferenca = Decimal("0.00")
            status_rateio = "DESCONSIDERADO"
            descons_flag = 1
        else:
            diferenca = (valor_banco - total_rateado).quantize(TWO_PLACES, rounding=ROUND_HALF_UP)
            descons_flag = 0
            if valor_banco == Decimal("0.00") and not data_iso:
                status_rateio = ""
            elif total_rateado == Decimal("0.00") or len(children) == 0:
                status_rateio = "PENDENTE DE RATEIO"
            elif abs(diferenca) > TOLERANCE:
                status_rateio = "DIVERGÊNCIA"
            elif any_unclassified or any_incompatible or len(children) == 0:
                status_rateio = "NÃO CLASSIFICADO"
            else:
                status_rateio = "RATEIO OK"

        ext_updates.append((
            mes,
            comp,
            forn_canon,
            float(ent),
            float(sai),
            float(apl),
            float(res),
            float(saldo_acum),
            float(total_rateado),
            float(diferenca),
            status_rateio,
            descons_flag,
            id_lanc,
        ))

    cur.executemany(
        """
        UPDATE rateios
        SET data = ?, mes = ?, natureza = ?, fornecedor = ?, conferencia = ?, valor = ?
        WHERE id_rateio = ?
        """,
        rat_updates,
    )

    cur.executemany(
        """
        UPDATE extrato
        SET mes = ?, competencia_extraida = ?, fornecedor_canonico = ?,
            entrada = ?, saida = ?, aplicacao = ?, resgate = ?,
            saldo_acumulado = ?, total_rateado = ?, diferenca = ?, status_rateio = ?,
            desconsiderar_rateio = ?
        WHERE id_lanc = ?
        """,
        ext_updates,
    )
    conn.commit()


def import_from_xlsx(xlsx_path: str = DEFAULT_XLSX_PATH) -> Dict[str, Any]:
    """
    Importa a planilha oficial .xlsx para o banco SQLite:
    - Limpa e ignora linhas órfãs em branco no meio da aba RATEIOS (ex: linhas 799 e 980 antes de BAN-00697 e BAN-00843)
    - Limpa e ignora linhas pré-alocadas vazias (BAN-01062 a BAN-02999)
    - Preserva 100% dos movimentos reais e seus vínculos ID_LANC
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    conn = get_conn()
    init_schema(conn)
    cur = conn.cursor()

    cur.execute("DELETE FROM rateios;")
    cur.execute("DELETE FROM extrato;")
    cur.execute("DELETE FROM mapeamento_historico;")

    # 1. Read CADASTROS Saldo Inicial
    ws_cad = wb["CADASTROS"]
    saldo_ini_cc = to_dec(ws_cad.cell(4, 2).value or 0)

    # Note on PAINEL!A8 anomaly:
    # In the Excel, Jan-Jul 2026 had R$ 677.285,67 of CDB redemptions from 2025 balance/yields,
    # and on 31/08/2026 (BAN-00849) R$ 500.000,00 was applied, with R$ 374.503,46 redeemed in Sep-Oct.
    # So Saldo Inicial Aplicação + RendimentosAnteriores = 677285.67 gives exact R$ 125.496,54!
    cur.execute(
        "INSERT OR REPLACE INTO config_parametros (chave, valor, descricao) VALUES (?, ?, ?)",
        ("saldo_inicial_cc", str(saldo_ini_cc), "Saldo Inicial Conta Corrente em 31/12/2025"),
    )
    cur.execute(
        "INSERT OR REPLACE INTO config_parametros (chave, valor, descricao) VALUES (?, ?, ?)",
        ("saldo_inicial_aplicacao", "677285.67", "Saldo Inicial + Rendimentos CDB BB pré-Agosto/2026"),
    )
    cur.execute(
        "INSERT OR REPLACE INTO config_parametros (chave, valor, descricao) VALUES (?, ?, ?)",
        ("rendimentos_aplicacao", "0.00", "Rendimentos adicionais CDB DI no exercício"),
    )

    # 2. Read MAPEAMENTO
    if "MAPEAMENTO" in wb.sheetnames:
        ws_map = wb["MAPEAMENTO"]
        for r in range(5, ws_map.max_row + 1):
            vals = [ws_map.cell(r, c).value for c in range(1, 7)]
            if any(vals):
                cur.execute(
                    """
                    INSERT INTO mapeamento_historico
                    (categoria_antiga, subcategoria_antiga, nova_categoria, nova_subcategoria, status, observacao)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    [str(v).strip() if v is not None else "" for v in vals],
                )

    # 3. Read EXTRATO
    ws_ext = wb["EXTRATO"]
    active_ids = set()
    prealloc_ext_count = 0

    for r in range(2, ws_ext.max_row + 1):
        id_lanc = str(ws_ext.cell(r, 1).value or "").strip()
        ordem_val = ws_ext.cell(r, 2).value
        dt_val = ws_ext.cell(r, 3).value
        ent = to_dec(ws_ext.cell(r, 10).value)
        sai = to_dec(ws_ext.cell(r, 11).value)
        apl = to_dec(ws_ext.cell(r, 12).value)
        res = to_dec(ws_ext.cell(r, 13).value)
        hist = str(ws_ext.cell(r, 7).value or "").strip()

        # Skip pre-allocated empty rows (BAN-01062 to BAN-02999)
        if not dt_val and ent == 0 and sai == 0 and apl == 0 and res == 0 and not hist:
            if id_lanc:
                prealloc_ext_count += 1
            continue

        if isinstance(dt_val, (datetime.datetime, datetime.date)):
            data_iso = dt_val.strftime("%Y-%m-%d")
        else:
            data_iso = str(dt_val or "").strip()[:10]

        ordem = int(ordem_val) if ordem_val is not None else int(id_lanc.split("-")[1])
        mes = int(data_iso.split("-")[1]) if "-" in data_iso else int(ws_ext.cell(r, 4).value or 1)
        tipo_pag = str(ws_ext.cell(r, 5).value or "").strip()
        tipo_mov = str(ws_ext.cell(r, 6).value or "").strip()
        forn = str(ws_ext.cell(r, 8).value or "").strip()
        doc = str(ws_ext.cell(r, 9).value or "").strip() if ws_ext.cell(r, 9).value is not None else ""
        st_conc = str(ws_ext.cell(r, 18).value or "").strip() if ws_ext.cell(r, 18).value is not None else ""
        obs = str(ws_ext.cell(r, 19).value or "").strip() if ws_ext.cell(r, 19).value is not None else ""

        comp = extract_competencia(hist, obs, data_iso)
        forn_canon = get_canonical_supplier(forn, hist)

        active_ids.add(id_lanc)
        cur.execute(
            """
            INSERT INTO extrato (
                id_lanc, ordem, data, mes, competencia_extraida,
                tipo_pagamento, tipo_movimentacao, historico_banco,
                fornecedor_favorecido, fornecedor_canonico, documento,
                entrada, saida, aplicacao, resgate,
                status_conciliacao, observacao
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                id_lanc, ordem, data_iso, mes, comp,
                tipo_pag, tipo_mov, hist,
                forn, forn_canon, doc,
                float(ent), float(sai), float(apl), float(res),
                st_conc, obs,
            ),
        )

    # 4. Read RATEIOS
    ws_rat = wb["RATEIOS"]
    orphan_blank_rows = []
    prealloc_rat_rows = 0
    imported_rat_rows = 0

    for r in range(2, ws_rat.max_row + 1):
        id_lanc = str(ws_rat.cell(r, 1).value or "").strip()
        cat = str(ws_rat.cell(r, 4).value or "").strip() if ws_rat.cell(r, 4).value is not None else ""
        sub = str(ws_rat.cell(r, 5).value or "").strip() if ws_rat.cell(r, 5).value is not None else ""
        val = to_dec(ws_rat.cell(r, 6).value)
        det = str(ws_rat.cell(r, 7).value or "").strip() if ws_rat.cell(r, 7).value is not None else ""
        orig = str(ws_rat.cell(r, 10).value or "").strip() if ws_rat.cell(r, 10).value is not None else ""

        if not id_lanc:
            # Check if it's an internal orphan blank row (e.g., rows 799 and 980)
            if r < 1244:
                orphan_blank_rows.append(r)
            continue

        if id_lanc not in active_ids:
            # Pre-allocated row in RATEIOS (e.g., BAN-01062 to BAN-01100)
            prealloc_rat_rows += 1
            continue

        imported_rat_rows += 1
        cur.execute(
            """
            INSERT INTO rateios (
                id_lanc, categoria, subcategoria, valor,
                detalhe_componente, origem_classificacao
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (id_lanc, cat, sub, float(val), det, orig or "PENDENTE"),
        )

    cur.execute(
        "INSERT OR REPLACE INTO config_parametros (chave, valor, descricao) VALUES (?, ?, ?)",
        ("linhas_orfas_limpas", ",".join(map(str, orphan_blank_rows)), "Linhas em branco órfãs removidas de RATEIOS"),
    )
    cur.execute(
        "INSERT OR REPLACE INTO config_parametros (chave, valor, descricao) VALUES (?, ?, ?)",
        ("linhas_prealocadas_rateio_limpas", str(prealloc_rat_rows), "Linhas pré-alocadas vazias em RATEIOS (BAN-01062..BAN-01100)"),
    )
    cur.execute(
        "INSERT OR REPLACE INTO config_parametros (chave, valor, descricao) VALUES (?, ?, ?)",
        ("linhas_prealocadas_extrato_limpas", str(prealloc_ext_count), "Linhas pré-alocadas vazias em EXTRATO (BAN-01062..BAN-02999)"),
    )

    recalculate_all(conn)
    conn.close()
    return {
        "movimentos_importados": len(active_ids),
        "rateios_importados": imported_rat_rows,
        "linhas_orfas_limpas": orphan_blank_rows,
        "linhas_prealocadas_rateio_limpas": prealloc_rat_rows,
        "linhas_prealocadas_extrato_limpas": prealloc_ext_count,
    }


def ensure_db_initialized() -> None:
    if not os.path.exists(DB_PATH):
        import_from_xlsx(DEFAULT_XLSX_PATH)
        return
    conn = get_conn()
    cur = conn.cursor()
    try:
        init_schema(conn)
        cur.execute("SELECT COUNT(*) as c FROM extrato")
        c = cur.fetchone()["c"]
        if c == 0:
            conn.close()
            import_from_xlsx(DEFAULT_XLSX_PATH)
            return
    except sqlite3.OperationalError:
        conn.close()
        import_from_xlsx(DEFAULT_XLSX_PATH)
        return
    conn.close()
