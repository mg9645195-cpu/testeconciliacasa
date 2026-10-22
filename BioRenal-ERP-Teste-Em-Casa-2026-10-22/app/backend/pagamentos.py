# -*- coding: utf-8 -*-
"""Programação de pagamentos e controle de prestadores no SQLite local."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import re
import sqlite3
import unicodedata
from urllib.parse import urlparse

from openpyxl import load_workbook

from database import get_conn


def _cents(value):
    if value is None or value == "":
        return None
    return int((Decimal(str(value)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _money(value):
    return None if value is None else float(Decimal(value) / 100)


def _date(value):
    if not value:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()[:10]
    date.fromisoformat(text)
    return text


def _source_date(value):
    """Source cells may contain '-' or notes rather than an actual date."""
    try:
        return _date(value)
    except ValueError:
        return ""


def _text(value):
    return "" if value is None else str(value).strip()


def _norm(value):
    text = unicodedata.normalize("NFKD", _text(value).upper())
    return re.sub(r"\s+", " ", "".join(c for c in text if not unicodedata.combining(c))).strip()


def _url(value):
    value = _text(value)
    if not value:
        return ""
    parsed = urlparse(value)
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Use um link HTTPS válido para o documento.")
    return value


def ensure_payment_schema(conn):
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS programacao_pagamentos (
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      origem TEXT NOT NULL DEFAULT 'MANUAL', chave_origem TEXT UNIQUE,
      aba_origem TEXT NOT NULL DEFAULT '', linha_origem INTEGER,
      vencimento TEXT NOT NULL, tipo TEXT NOT NULL DEFAULT '', nf TEXT NOT NULL DEFAULT '',
      favorecido TEXT NOT NULL, valor_centavos INTEGER CHECK(valor_centavos >= 0),
      observacoes TEXT NOT NULL DEFAULT '', documento_url TEXT NOT NULL DEFAULT '',
      dda_status TEXT NOT NULL DEFAULT 'NAO_CONFERIDO'
        CHECK(dda_status IN ('NAO_CONFERIDO','CONFERIDO','DIVERGENTE','NAO_APLICA')),
      dda_valor_centavos INTEGER CHECK(dda_valor_centavos >= 0),
      dda_vencimento TEXT NOT NULL DEFAULT '',
      status TEXT NOT NULL DEFAULT 'PREVISTO'
        CHECK(status IN ('PREVISTO','ENVIADO','PAGO','CANCELADO')),
      pago_em TEXT NOT NULL DEFAULT '', anomalia TEXT NOT NULL DEFAULT '',
      atualizado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    );
    CREATE INDEX IF NOT EXISTS idx_programacao_vencimento ON programacao_pagamentos(vencimento);
    CREATE TABLE IF NOT EXISTS programacao_auditoria (
      id INTEGER PRIMARY KEY AUTOINCREMENT, item_id INTEGER NOT NULL,
      acao TEXT NOT NULL, anterior TEXT NOT NULL DEFAULT '', posterior TEXT NOT NULL DEFAULT '',
      criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS documentos_extrato (
      id_lanc TEXT PRIMARY KEY, data TEXT NOT NULL DEFAULT '',
      historico TEXT NOT NULL DEFAULT '', url TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS cadastro_bancario_prestadores (
      id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT NOT NULL UNIQUE,
      tipo TEXT NOT NULL DEFAULT '', banco TEXT NOT NULL DEFAULT '', agencia TEXT NOT NULL DEFAULT '',
      conta TEXT NOT NULL DEFAULT '', documento TEXT NOT NULL DEFAULT '', pix TEXT NOT NULL DEFAULT '',
      observacoes TEXT NOT NULL DEFAULT '', conferido INTEGER NOT NULL DEFAULT 0,
      atualizado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS pagamentos_prestadores (
      id INTEGER PRIMARY KEY AUTOINCREMENT, origem TEXT NOT NULL DEFAULT 'MANUAL',
      chave_origem TEXT UNIQUE, tipo TEXT NOT NULL CHECK(tipo IN ('PRESTADOR','PLANTAO')),
      competencia TEXT NOT NULL DEFAULT '', nf TEXT NOT NULL DEFAULT '',
      nome TEXT NOT NULL, descricao TEXT NOT NULL DEFAULT '',
      quantidade REAL, valor_unitario_centavos INTEGER,
      valor_centavos INTEGER CHECK(valor_centavos >= 0), vencimento TEXT NOT NULL DEFAULT '',
      enviado INTEGER NOT NULL DEFAULT 0, pago_em TEXT NOT NULL DEFAULT '',
      observacoes TEXT NOT NULL DEFAULT '', banco_id INTEGER,
      anomalia TEXT NOT NULL DEFAULT '', atualizado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
      FOREIGN KEY(banco_id) REFERENCES cadastro_bancario_prestadores(id)
    );
    CREATE INDEX IF NOT EXISTS idx_prestadores_vencimento ON pagamentos_prestadores(vencimento);
    """)
    columns = {row["name"] for row in conn.execute("PRAGMA table_info(programacao_pagamentos)")}
    if "extrato_id_lanc" not in columns:
        conn.execute("ALTER TABLE programacao_pagamentos ADD COLUMN extrato_id_lanc TEXT NOT NULL DEFAULT ''")
    if "vinculado_em" not in columns:
        conn.execute("ALTER TABLE programacao_pagamentos ADD COLUMN vinculado_em TEXT NOT NULL DEFAULT ''")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_programacao_extrato_vinculo ON programacao_pagamentos(extrato_id_lanc) WHERE extrato_id_lanc != ''")
    conn.commit()


def _rows(conn, sql, args=()):
    return [dict(row) for row in conn.execute(sql, args).fetchall()]


def get_programacao(inicio="", fim=""):
    conn = get_conn()
    ensure_payment_schema(conn)
    args = []
    where = "WHERE 1=1"
    if inicio:
        where += " AND vencimento >= ?"
        args.append(_date(inicio))
    if fim:
        where += " AND vencimento <= ?"
        args.append(_date(fim))
    rows = _rows(conn, f"""SELECT p.*, e.status_rateio AS extrato_status_rateio,
        e.data AS extrato_data FROM programacao_pagamentos p
        LEFT JOIN extrato e ON e.id_lanc=p.extrato_id_lanc
        {where} ORDER BY p.vencimento,p.id""", args)
    links = _rows(conn, "SELECT * FROM documentos_extrato ORDER BY data DESC,id_lanc DESC")
    conn.close()
    for row in rows:
        row["valor"] = _money(row.pop("valor_centavos"))
        row["dda_valor"] = _money(row.pop("dda_valor_centavos"))
    return {"itens": rows, "documentos_extrato": links}


def save_programacao(body):
    conn = get_conn()
    ensure_payment_schema(conn)
    item_id = body.get("id")
    old = conn.execute("SELECT * FROM programacao_pagamentos WHERE id=?", (item_id,)).fetchone() if item_id else None
    if item_id and not old:
        raise ValueError("Lançamento de programação não encontrado.")
    vencimento = _date(body.get("vencimento"))
    favorecido = _text(body.get("favorecido"))
    if not vencimento or not favorecido:
        raise ValueError("Informe vencimento e favorecido.")
    valor = _cents(body.get("valor"))
    dda_valor = _cents(body.get("dda_valor"))
    if (valor is not None and valor < 0) or (dda_valor is not None and dda_valor < 0):
        raise ValueError("Valores não podem ser negativos.")
    dda = body.get("dda_status", "NAO_CONFERIDO")
    status = body.get("status", "PREVISTO")
    if dda not in ("NAO_CONFERIDO", "CONFERIDO", "DIVERGENTE", "NAO_APLICA") or status not in ("PREVISTO", "ENVIADO", "PAGO", "CANCELADO"):
        raise ValueError("Status inválido.")
    pago_em = _date(body.get("pago_em"))
    if status == "PAGO" and not pago_em:
        raise ValueError("Informe a data efetiva do pagamento.")
    if old and old["extrato_id_lanc"] and (status != "PAGO" or valor != old["valor_centavos"]):
        raise ValueError("Desvincule o movimento do extrato antes de alterar a situação ou o valor.")
    dda_vencimento = _date(body.get("dda_vencimento"))
    if dda == "CONFERIDO" and (valor is None or dda_valor is None or valor != dda_valor or (dda_vencimento and vencimento != dda_vencimento)):
        raise ValueError("DDA conferido exige valor igual e vencimento compatível.")
    vals = (vencimento, _text(body.get("tipo")), _text(body.get("nf")), favorecido, valor,
            _text(body.get("observacoes")), _url(body.get("documento_url")), dda,
            dda_valor, dda_vencimento, status, pago_em)
    with conn:
        if old:
            conn.execute("""UPDATE programacao_pagamentos SET vencimento=?,tipo=?,nf=?,favorecido=?,valor_centavos=?,
              observacoes=?,documento_url=?,dda_status=?,dda_valor_centavos=?,dda_vencimento=?,status=?,pago_em=?,
              atualizado_em=CURRENT_TIMESTAMP WHERE id=?""", vals + (item_id,))
        else:
            item_id = conn.execute("""INSERT INTO programacao_pagamentos
              (vencimento,tipo,nf,favorecido,valor_centavos,observacoes,documento_url,dda_status,
              dda_valor_centavos,dda_vencimento,status,pago_em) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""", vals).lastrowid
        conn.execute("INSERT INTO programacao_auditoria(item_id,acao,anterior,posterior) VALUES (?,?,?,?)",
                     (item_id, "ATUALIZAR" if old else "CRIAR", str(dict(old)) if old else "", str(vals)))
    conn.close()
    return {"sucesso": True, "id": item_id}


def get_candidatos_conciliacao(item_id):
    conn = get_conn()
    ensure_payment_schema(conn)
    item = conn.execute("SELECT * FROM programacao_pagamentos WHERE id=?", (item_id,)).fetchone()
    if not item:
        conn.close()
        raise ValueError("Gasto previsto não encontrado.")
    if item["status"] != "PAGO" or item["valor_centavos"] is None:
        conn.close()
        raise ValueError("Confirme o valor e a data do pagamento antes de vincular ao extrato.")
    paid = date.fromisoformat(item["pago_em"])
    candidates = []
    for row in conn.execute("""SELECT e.id_lanc,e.data,e.historico_banco,e.fornecedor_canonico,
        e.documento,e.saida,e.status_rateio,p.id AS usado_por
        FROM extrato e LEFT JOIN programacao_pagamentos p ON p.extrato_id_lanc=e.id_lanc
        WHERE e.saida>0 AND COALESCE(e.desconsiderar_rateio,0)=0"""):
        if _cents(row["saida"]) != item["valor_centavos"]:
            continue
        entry = dict(row)
        if entry["usado_por"] and entry["usado_por"] != item_id:
            continue
        entry["dias_distancia"] = abs((date.fromisoformat(entry["data"]) - paid).days)
        entry["valor"] = _money(item["valor_centavos"])
        entry.pop("saida")
        candidates.append(entry)
    conn.close()
    candidates.sort(key=lambda row: (row["dias_distancia"],row["data"],row["id_lanc"]))
    return {"item_id": item_id, "vinculado": item["extrato_id_lanc"], "candidatos": candidates[:60]}


def vincular_programacao_extrato(body):
    item_id = int(body.get("id") or 0)
    id_lanc = _text(body.get("id_lanc")).upper()
    conn = get_conn()
    ensure_payment_schema(conn)
    with conn:
        item = conn.execute("SELECT * FROM programacao_pagamentos WHERE id=?", (item_id,)).fetchone()
        if not item:
            raise ValueError("Gasto previsto não encontrado.")
        previous = item["extrato_id_lanc"]
        if not id_lanc:
            if not previous:
                raise ValueError("O gasto ainda não está vinculado ao extrato.")
            conn.execute("UPDATE programacao_pagamentos SET extrato_id_lanc='',vinculado_em='',atualizado_em=CURRENT_TIMESTAMP WHERE id=?", (item_id,))
            action = "DESVINCULAR_EXTRATO"
        else:
            if item["status"] != "PAGO" or not item["pago_em"] or item["valor_centavos"] is None:
                raise ValueError("Confirme valor e pagamento antes da conciliação.")
            movement = conn.execute("SELECT * FROM extrato WHERE id_lanc=?", (id_lanc,)).fetchone()
            if not movement or movement["tipo_movimentacao"] != "SAÍDA" or movement["desconsiderar_rateio"]:
                raise ValueError("Selecione uma saída válida do extrato.")
            if _cents(movement["saida"]) != item["valor_centavos"]:
                raise ValueError("O valor da saída bancária difere do gasto confirmado.")
            used = conn.execute("SELECT id FROM programacao_pagamentos WHERE extrato_id_lanc=? AND id!=?", (id_lanc,item_id)).fetchone()
            if used:
                raise ValueError("Este movimento bancário já está vinculado a outro gasto.")
            conn.execute("UPDATE programacao_pagamentos SET extrato_id_lanc=?,vinculado_em=CURRENT_TIMESTAMP,atualizado_em=CURRENT_TIMESTAMP WHERE id=?", (id_lanc,item_id))
            action = "VINCULAR_EXTRATO"
        conn.execute("INSERT INTO programacao_auditoria(item_id,acao,anterior,posterior) VALUES (?,?,?,?)", (item_id,action,previous,id_lanc))
    conn.close()
    return {"sucesso":True,"id":item_id,"extrato_id_lanc":id_lanc,
            "status_rateio":movement["status_rateio"] if id_lanc else ""}


def import_programacao_xlsx(data):
    wb = load_workbook(BytesIO(data), data_only=False)
    conn = get_conn()
    ensure_payment_schema(conn)
    added = 0
    anomalies = 0
    with conn:
        for ws in wb:
            if not re.fullmatch(r"(?:JAN|FEV|MAR|ABR|MAI|JUN|JUL|AGO|SET|OUT|NOV|DEZ)-\d{2}", ws.title.upper()):
                continue
            month_code, year_code = ws.title.upper().split("-")
            expected_month = {"JAN": 1,"FEV": 2,"MAR": 3,"ABR": 4,"MAI": 5,"JUN": 6,
                              "JUL": 7,"AGO": 8,"SET": 9,"OUT": 10,"NOV": 11,"DEZ": 12}[month_code]
            expected_year = 2000 + int(year_code)
            for row in ws.iter_rows(min_row=2):
                if not isinstance(row[0].value, (date, datetime)):
                    continue
                due = _date(row[0].value)
                payee = _text(row[3].value)
                if not payee:
                    continue
                amount = _cents(row[4].value) if isinstance(row[4].value, (int, float, Decimal)) and row[4].value > 0 else None
                issue = "DATA_FORA_DA_ABA" if (row[0].value.year, row[0].value.month) != (expected_year, expected_month) else ""
                if issue:
                    anomalies += 1
                link = next((c.hyperlink.target for c in row if c.hyperlink and c.hyperlink.target), "")
                key = f"PROJECAO:{ws.title}:{row[0].row}"
                result = conn.execute("""INSERT OR IGNORE INTO programacao_pagamentos
                  (origem,chave_origem,aba_origem,linha_origem,vencimento,tipo,nf,favorecido,valor_centavos,
                   observacoes,documento_url,anomalia) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                  ("PLANILHA_PROJECAO",key,ws.title,row[0].row,due,_text(row[1].value),
                   _text(row[2].value),payee,amount,_text(row[5].value),_url(link) if link else "",issue))
                added += result.rowcount
    conn.close()
    return {"incluidos": added, "datas_fora_da_aba": anomalies}


def import_extrato_links(data):
    wb = load_workbook(BytesIO(data), data_only=False)
    ws = wb["EXTRATO"]
    conn = get_conn()
    ensure_payment_schema(conn)
    count = 0
    with conn:
        for row in ws.iter_rows(min_row=2):
            lanc = _text(row[0].value)
            cell = row[6]
            if not re.fullmatch(r"BAN-\d{5}", lanc) or not cell.hyperlink or not cell.hyperlink.target:
                continue
            url = _url(cell.hyperlink.target)
            conn.execute("INSERT OR IGNORE INTO documentos_extrato(id_lanc,data,historico,url) VALUES (?,?,?,?)",
                         (lanc,_date(row[2].value),_text(cell.value),url))
            count += 1
    conn.close()
    return {"links_identificados": count}


def import_prestadores_xlsx(data):
    wb = load_workbook(BytesIO(data), data_only=True)
    conn = get_conn()
    ensure_payment_schema(conn)
    bank_added = payment_added = zero_amount = 0
    with conn:
        bank_ws = next(s for s in wb if _norm(s.title) == "CADASTRO BANCARIO")
        for row in bank_ws.iter_rows(min_row=3, values_only=True):
            name = _text(row[0])
            if not name:
                continue
            result = conn.execute("""INSERT OR IGNORE INTO cadastro_bancario_prestadores
              (nome,tipo,banco,agencia,conta,documento,pix,observacoes) VALUES (?,?,?,?,?,?,?,?)""",
              (name,*[_text(row[i]) for i in range(1,8)]))
            bank_added += result.rowcount
        profiles = {_norm(r["nome"]): r["id"] for r in _rows(conn,"SELECT id,nome FROM cadastro_bancario_prestadores")}
        for ws in wb:
            title = _norm(ws.title)
            if title not in ("PRESTADORES E FORNECEDORES", "PLANTOES MEDICOS"):
                continue
            medical = title == "PLANTOES MEDICOS"
            for cells in ws.iter_rows(min_row=3, values_only=False):
                v = [c.value for c in cells]
                if not _norm(v[0]).startswith(("REALIZADO", "PROJECAO")):
                    continue
                name = _text(v[3])
                if not name:
                    continue
                total = v[6] if medical else v[5]
                amount = _cents(total) if isinstance(total,(int,float,Decimal)) and total > 0 else None
                if amount is None:
                    zero_amount += 1
                due = _source_date(v[7] if medical else v[6])
                paid = _source_date(v[9] if medical else v[8])
                sent = _norm(v[8] if medical else v[7]) in ("SIM", "S")
                bank_id = profiles.get(_norm(name))  # Only exact canonical match: never guess a beneficiary.
                key = f"PRESTADORES:{ws.title}:{cells[0].row}"
                issue = "DADOS_BANCARIOS_A_CONFERIR" if bank_id is None else ""
                quantity = v[4] if medical and isinstance(v[4],(int,float)) else None
                unit = _cents(v[5]) if medical and isinstance(v[5],(int,float)) and v[5] > 0 else None
                result = conn.execute("""INSERT OR IGNORE INTO pagamentos_prestadores
                  (origem,chave_origem,tipo,competencia,nf,nome,descricao,quantidade,valor_unitario_centavos,
                   valor_centavos,vencimento,enviado,pago_em,observacoes,banco_id,anomalia)
                  VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                  ("PLANILHA_PRESTADORES",key,"PLANTAO" if medical else "PRESTADOR",
                   _source_date(v[1])[:7],_text(v[2]),name,"Plantões médicos" if medical else _text(v[4]),
                   quantity,unit,amount,due,int(sent),paid,_text(v[12] if medical else v[11]),bank_id,issue))
                payment_added += result.rowcount
    conn.close()
    return {"prestadores_incluidos": payment_added, "cadastros_bancarios_incluidos": bank_added,
            "valores_a_confirmar": zero_amount}


def get_prestadores(inicio="", fim="", include_bank=False):
    conn = get_conn()
    ensure_payment_schema(conn)
    args = []
    where = "WHERE 1=1"
    if inicio:
        where += " AND p.vencimento >= ?"
        args.append(_date(inicio))
    if fim:
        where += " AND p.vencimento <= ?"
        args.append(_date(fim))
    rows = _rows(conn,"""SELECT p.*,b.nome AS banco_nome,b.conferido AS banco_conferido""" +
                 (",b.banco,b.agencia,b.conta,b.documento,b.pix,b.observacoes AS banco_observacoes" if include_bank else "") +
                 " FROM pagamentos_prestadores p LEFT JOIN cadastro_bancario_prestadores b ON b.id=p.banco_id " + where +
                 " ORDER BY p.vencimento,p.nome,p.id",args)
    profiles = _rows(conn,"SELECT id,nome,tipo,conferido FROM cadastro_bancario_prestadores ORDER BY nome")
    conn.close()
    today = date.today().isoformat()
    for row in rows:
        row["valor"] = _money(row.pop("valor_centavos"))
        row["valor_unitario"] = _money(row.pop("valor_unitario_centavos"))
        row["situacao"] = "PAGO" if row["pago_em"] else ("A_CONFIRMAR" if row["valor"] is None else
                         "ATRASADO" if row["vencimento"] and row["vencimento"] < today else "PENDENTE")
    return {"itens": rows, "cadastros": profiles}


def save_prestador(body):
    conn = get_conn()
    ensure_payment_schema(conn)
    item_id = body.get("id")
    old = conn.execute("SELECT id FROM pagamentos_prestadores WHERE id=?",(item_id,)).fetchone() if item_id else None
    if item_id and not old:
        raise ValueError("Pagamento não encontrado.")
    typ = body.get("tipo", "PRESTADOR")
    if typ not in ("PRESTADOR", "PLANTAO"):
        raise ValueError("Tipo inválido.")
    name = _text(body.get("nome"))
    if not name:
        raise ValueError("Informe o prestador ou médico.")
    val = _cents(body.get("valor"))
    unit = _cents(body.get("valor_unitario"))
    qty = body.get("quantidade")
    qty = float(qty) if qty not in (None, "") else None
    if val is not None and val < 0 or qty is not None and qty < 0:
        raise ValueError("Valores não podem ser negativos.")
    if qty == 0:
        qty = None
    if typ == "PLANTAO" and qty is not None and unit is not None:
        calculated = int((Decimal(str(qty)) * Decimal(unit)).quantize(Decimal("1"),rounding=ROUND_HALF_UP))
        if val is None:
            val = calculated
        elif val != calculated:
            raise ValueError("Valor total deve corresponder à quantidade × valor por plantão.")
    bank_id = body.get("banco_id") or None
    if bank_id and not conn.execute("SELECT 1 FROM cadastro_bancario_prestadores WHERE id=?",(bank_id,)).fetchone():
        raise ValueError("Cadastro bancário inexistente.")
    paid = _date(body.get("pago_em"))
    vals = (typ,_text(body.get("competencia")),_text(body.get("nf")),name,_text(body.get("descricao")),
            qty,unit,val,_date(body.get("vencimento")),int(bool(body.get("enviado"))),paid,
            _text(body.get("observacoes")),bank_id)
    with conn:
        if old:
            conn.execute("""UPDATE pagamentos_prestadores SET tipo=?,competencia=?,nf=?,nome=?,descricao=?,
              quantidade=?,valor_unitario_centavos=?,valor_centavos=?,vencimento=?,enviado=?,pago_em=?,
              observacoes=?,banco_id=?,atualizado_em=CURRENT_TIMESTAMP WHERE id=?""",vals+(item_id,))
        else:
            item_id = conn.execute("""INSERT INTO pagamentos_prestadores
              (tipo,competencia,nf,nome,descricao,quantidade,valor_unitario_centavos,valor_centavos,
               vencimento,enviado,pago_em,observacoes,banco_id) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",vals).lastrowid
    conn.close()
    return {"sucesso":True,"id":item_id}


def save_banco(body):
    conn = get_conn()
    ensure_payment_schema(conn)
    bank_id = body.get("id")
    name = _text(body.get("nome"))
    if not name:
        raise ValueError("Informe o nome do favorecido.")
    vals = (name,*[_text(body.get(k)) for k in ("tipo","banco","agencia","conta","documento","pix","observacoes")],
            int(bool(body.get("conferido"))))
    with conn:
        if bank_id:
            cur = conn.execute("""UPDATE cadastro_bancario_prestadores SET nome=?,tipo=?,banco=?,agencia=?,
              conta=?,documento=?,pix=?,observacoes=?,conferido=?,atualizado_em=CURRENT_TIMESTAMP WHERE id=?""",
              vals+(bank_id,))
            if not cur.rowcount:
                raise ValueError("Cadastro bancário não encontrado.")
        else:
            bank_id = conn.execute("""INSERT INTO cadastro_bancario_prestadores
              (nome,tipo,banco,agencia,conta,documento,pix,observacoes,conferido)
              VALUES (?,?,?,?,?,?,?,?,?)""",vals).lastrowid
    conn.close()
    return {"sucesso":True,"id":bank_id}


def get_banco(bank_id):
    conn = get_conn()
    ensure_payment_schema(conn)
    row = conn.execute("SELECT * FROM cadastro_bancario_prestadores WHERE id=?",(bank_id,)).fetchone()
    conn.close()
    if not row:
        raise ValueError("Cadastro bancário não encontrado.")
    return dict(row)


def marcar_prestadores_enviados(body):
    raw_ids = body.get("ids", [])
    if not isinstance(raw_ids, list) or not raw_ids or len(raw_ids) > 200:
        raise ValueError("Selecione até 200 lançamentos para registrar o envio.")
    ids = [int(value) for value in raw_ids]
    if len(ids) != len(set(ids)):
        raise ValueError("Há lançamentos repetidos na seleção.")
    conn = get_conn()
    ensure_payment_schema(conn)
    with conn:
        for item_id in ids:
            row = conn.execute("""SELECT p.valor_centavos,p.pago_em,p.enviado,b.conferido AS banco_conferido
                FROM pagamentos_prestadores p
                LEFT JOIN cadastro_bancario_prestadores b ON b.id=p.banco_id
                WHERE p.id=?""",(item_id,)).fetchone()
            if not row or row["valor_centavos"] is None or row["pago_em"] or not row["banco_conferido"]:
                raise ValueError(f"Revise valor, situação e cadastro bancário do lançamento {item_id} antes de registrar o envio.")
        conn.executemany("UPDATE pagamentos_prestadores SET enviado=1,atualizado_em=CURRENT_TIMESTAMP WHERE id=?", [(item_id,) for item_id in ids])
    conn.close()
    return {"sucesso":True,"enviados":len(ids)}
